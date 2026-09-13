"""Synthetic fixtures exercise extraction, source scope and local API persistence."""

import io
import os
import zipfile

os.environ["DATABASE_URL"] = "sqlite://"
os.environ["DEFAULT_API_KEY"] = ""
os.environ["MEMORY_RETRIEVAL_MODE"] = "lexical"

import pytest
from docx import Document
from fastapi import FastAPI, HTTPException
from fastapi.testclient import TestClient
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
from sqlmodel import Session, SQLModel, create_engine, select

from app import documents
from app.db import get_session
from app.documents import (
    DocumentAttachment,
    DocumentParseError,
    document_summary,
    parse_document,
    parse_document_limited,
    resolve_documents,
)
from app.models import KnowledgeTree, Message, Node


def pdf_bytes(texts=("A validation set checks generalization.",), encrypted=False):
    writer = PdfWriter()
    for text in texts:
        page = writer.add_blank_page(width=300, height=300)
        if text:
            font = DictionaryObject(
                {
                    NameObject("/Type"): NameObject("/Font"),
                    NameObject("/Subtype"): NameObject("/Type1"),
                    NameObject("/BaseFont"): NameObject("/Helvetica"),
                }
            )
            page[NameObject("/Resources")] = DictionaryObject(
                {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})}
            )
            stream = DecodedStreamObject()
            stream.set_data(f"BT /F1 12 Tf 20 250 Td ({text}) Tj ET".encode())
            page[NameObject("/Contents")] = stream
    if encrypted:
        writer.encrypt("synthetic-test-password", algorithm="RC4-128")
    target = io.BytesIO()
    writer.write(target)
    return target.getvalue()


def docx_bytes():
    document = Document()
    document.add_paragraph("学习笔记：验证集用于检查泛化能力。")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Split"
    table.cell(0, 1).text = "Purpose"
    table.cell(1, 0).text = "Validation"
    table.cell(1, 1).text = "Model selection"
    document.add_paragraph("Keep the test set separate.")
    target = io.BytesIO()
    document.save(target)
    return target.getvalue()


@pytest.fixture
def setup(tmp_path, monkeypatch):
    engine = create_engine(
        f"sqlite:///{tmp_path / 'documents.db'}", connect_args={"check_same_thread": False}
    )
    SQLModel.metadata.create_all(engine)
    app = FastAPI()
    app.include_router(documents.router)

    def session_override():
        with Session(engine) as session:
            yield session

    app.dependency_overrides[get_session] = session_override
    # Extraction rules are tested directly; one test below covers the real subprocess.
    monkeypatch.setattr(documents, "parse_document_limited", parse_document)
    with Session(engine) as session:
        tree = KnowledgeTree(title="Document learning")
        other_tree = KnowledgeTree(title="Unrelated learning")
        session.add_all([tree, other_tree])
        session.flush()
        node = Node(tree_id=tree.id, title="Read source")
        other_node = Node(tree_id=other_tree.id, title="Other source")
        session.add_all([node, other_node])
        session.commit()
        ids = node.id, other_node.id, tree.id, other_tree.id
    with TestClient(app) as client:
        yield client, engine, ids
    engine.dispose()


def upload(client, node_id, name="notes.txt", content=b"Test source text"):
    return client.post(
        f"/nodes/{node_id}/documents", files={"file": (name, content, "application/octet-stream")}
    )


def test_pdf_preserves_page_sources_and_warns_about_empty_pages():
    sections, warnings = parse_document(pdf_bytes(("First source.", "", "Third source.")), "a.pdf")
    assert sections == [
        {"label": "Page 1", "text": "First source."},
        {"label": "Page 3", "text": "Third source."},
    ]
    assert len(warnings) == 1 and "2" in warnings[0] and "OCR" in warnings[0]


def test_docx_preserves_paragraph_table_order():
    sections, warnings = parse_document(docx_bytes(), "notes.docx")
    assert [section["label"] for section in sections] == ["Paragraph 1", "Table 1", "Paragraph 2"]
    assert "验证集" in sections[0]["text"]
    assert sections[1]["text"] == "Split\tPurpose\nValidation\tModel selection"
    assert "test set separate" in sections[2]["text"]
    assert warnings == []


@pytest.mark.parametrize("extension", ["txt", "md", "csv", "tsv", "json", "log"])
def test_utf8_text_inputs(extension):
    sections, warnings = parse_document(
        "\ufeff学习原文\nsecond line".encode(), f"notes.{extension}"
    )
    assert sections == [{"label": "Text", "text": "学习原文\nsecond line"}]
    assert warnings == []


@pytest.mark.parametrize(
    "content,name,reason",
    [
        (b"old binary", "old.doc", "另存"),
        (b"executable", "program.exe", "支持"),
        (b"", "empty.txt", "为空"),
        (b"   \n", "empty.txt", "没有可提取"),
        (b"\xff\xfe", "legacy.txt", "UTF-8"),
        (b"hello\x00data", "binary.txt", "二进制"),
        (b"not a pdf", "fake.pdf", "有效"),
        (b"%PDF-1.4\ninvalid", "broken.pdf", "损坏"),
        (b"not a docx", "fake.docx", "有效"),
    ],
)
def test_invalid_inputs_have_actionable_errors(content, name, reason):
    with pytest.raises(DocumentParseError, match=reason):
        parse_document(content, name)


def test_pdf_scans_and_encryption_are_explicit():
    with pytest.raises(DocumentParseError, match="OCR"):
        parse_document(pdf_bytes(("",)), "scan.pdf")
    with pytest.raises(DocumentParseError, match="密码"):
        parse_document(pdf_bytes(encrypted=True), "encrypted.pdf")


def test_extraction_limits(monkeypatch):
    with pytest.raises(DocumentParseError, match="200,000"):
        parse_document(b"x" * (documents.MAX_EXTRACTED_CHARACTERS + 1), "long.txt")
    monkeypatch.setattr(documents, "MAX_FILE_BYTES", 3)
    with pytest.raises(DocumentParseError, match="10 MiB"):
        parse_document(b"four", "big.txt")
    monkeypatch.setattr(documents, "MAX_FILE_BYTES", 10 * 1024 * 1024)
    monkeypatch.setattr(documents, "MAX_PDF_PAGES", 2)
    with pytest.raises(DocumentParseError, match="200 页"):
        parse_document(pdf_bytes(("one", "two", "three")), "long.pdf")


def test_docx_rejects_expansion_and_xml_entities(monkeypatch):
    monkeypatch.setattr(documents, "MAX_ZIP_BYTES", 50)
    with pytest.raises(DocumentParseError, match="解压"):
        parse_document(docx_bytes(), "large.docx")
    monkeypatch.setattr(documents, "MAX_ZIP_BYTES", 25 * 1024 * 1024)
    target = io.BytesIO()
    with zipfile.ZipFile(target, "w") as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr("word/document.xml", '<!DOCTYPE doc [<!ENTITY test "bad">]><doc/>')
    with pytest.raises(DocumentParseError, match="XML"):
        parse_document(target.getvalue(), "entities.docx")


def test_real_subprocess_returns_sources_and_errors():
    sections, warnings = parse_document_limited(pdf_bytes(), "process.pdf")
    assert "generalization" in sections[0]["text"]
    assert warnings == []
    with pytest.raises(DocumentParseError, match="另存"):
        parse_document_limited(b"old", "old.doc")


def test_subprocess_timeout_is_bounded(monkeypatch):
    monkeypatch.setattr(documents, "PARSE_TIMEOUT_SECONDS", 0)
    with pytest.raises(DocumentParseError, match="超时"):
        parse_document_limited(b"bounded", "short.txt")


@pytest.mark.parametrize(
    "name,content", [("report.pdf", pdf_bytes()), ("notes.docx", docx_bytes())]
)
def test_upload_preview_download_persist_exact_source(setup, name, content):
    client, engine, (node_id, _, tree_id, _) = setup
    response = upload(client, node_id, name, content)
    assert response.status_code == 200, response.text
    summary = response.json()
    assert set(summary) == {"id", "name", "media_type", "size", "characters", "warnings"}
    assert summary["name"] == name and summary["size"] == len(content)
    assert summary["characters"] > 0
    preview = client.get(f"/documents/{summary['id']}")
    assert preview.status_code == 200 and preview.json()["sections"][0]["text"]
    download = client.get(f"/documents/{summary['id']}/download")
    assert download.content == content
    assert download.headers["content-disposition"].startswith("attachment;")
    assert download.headers["x-content-type-options"] == "nosniff"
    with Session(engine) as session:
        stored = session.get(DocumentAttachment, summary["id"])
        assert stored.content == content and stored.tree_id == tree_id
        assert document_summary(stored).model_dump() == summary


def test_repeated_upload_deduplicates_only_within_current_tree(setup):
    client, engine, (node_id, other_node_id, tree_id, _) = setup
    first = upload(client, node_id).json()
    repeated = upload(client, node_id).json()
    other = upload(client, other_node_id).json()
    assert repeated["id"] == first["id"] and other["id"] != first["id"]
    with Session(engine) as session:
        assert len(session.exec(select(DocumentAttachment)).all()) == 2
        assert len(resolve_documents(session, [first["id"], first["id"]], tree_id)) == 1
        with pytest.raises(HTTPException, match="当前学习主题"):
            resolve_documents(session, [other["id"]], tree_id)
        with pytest.raises(HTTPException, match="不存在"):
            resolve_documents(session, ["missing"], tree_id)
        with pytest.raises(HTTPException, match="最多"):
            resolve_documents(session, [str(i) for i in range(5)], tree_id)


def test_delete_only_unused_sources(setup):
    client, engine, (node_id, _, _, _) = setup
    unused = upload(client, node_id, "unused.txt").json()["id"]
    used = upload(client, node_id, "used.txt").json()["id"]
    with Session(engine) as session:
        session.add(Message(node_id=node_id, role="user", content="Read it", document_ids=[used]))
        session.commit()
    assert client.delete(f"/documents/{used}").status_code == 409
    assert client.delete(f"/documents/{unused}").status_code == 204
    assert client.get(f"/documents/{unused}").status_code == 404
    assert client.get(f"/documents/{used}/download").status_code == 200


def test_api_limits_invalid_nodes_and_safe_filenames(setup, monkeypatch):
    client, engine, (node_id, _, _, _) = setup
    assert upload(client, 99999).status_code == 404
    assert upload(client, node_id, "legacy.doc").status_code == 422
    safe = upload(client, node_id, "../../学习.txt").json()
    assert safe["name"] == "学习.txt"
    download = client.get(f"/documents/{safe['id']}/download")
    assert "filename*=UTF-8''%E5%AD%A6%E4%B9%A0.txt" in download.headers["content-disposition"]
    monkeypatch.setattr(documents, "MAX_FILE_BYTES", 10)
    assert upload(client, node_id, content=b"x" * 11).status_code == 413
    with Session(engine) as session:
        assert len(session.exec(select(DocumentAttachment)).all()) == 1
